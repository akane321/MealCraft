"""Per-field agreement of the reviewer's filled sample (review-table-20261024) with the imported values.

Only rows the reviewer marked 是 in "已亲自核对（填 是）" count; the others are reported as unchecked, and a checked row
with no value is reported, not counted. The rows must be the sample's (`sample.json`). Offline, with a Python that has
openpyxl, from the repository root:
`python -I -X utf8 data-engineering/scripts/review_table_accuracy.py <filled sample.xlsx> --out <accuracy.md>`.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from import_review_table import CHECKED, FIELDS, OUT, SAMPLE_SIZE, SKIP, YOUR_VALUE, code, shown  # noqa: E402

SAMPLE = OUT.parent / "sample.json"


def same(field: str, imported, value) -> bool:
    return (
        sorted(imported) == value if field == "protein_sources" else imported == value and type(imported) is type(value)
    )


def accuracy(rows: list[dict], sample: dict) -> dict:
    """Per field: the sample's rows, the checked ones, how many agree, and the unchecked ones. `rows` are the
    filled sheet's `{"id", "field", "value", "checked"}`."""
    imported = {
        (recipe["id"], field): value for recipe in sample["recipes"] for field, value in recipe["imported"].items()
    }
    seen = {(row["id"], row["field"]) for row in rows}
    if seen != set(imported) or len(rows) != len(imported):
        raise SystemExit(f"the filled sheet's rows are not the sample's: {sorted(seen ^ set(imported))[:5]}")
    fields = {
        field: {"rows": 0, "checked": 0, "agree": 0, "unchecked": 0, "checked_without_value": 0} for field in FIELDS
    }
    differ = []
    for row in rows:
        counts = fields[row["field"]]
        counts["rows"] += 1
        value = code(row["value"])
        if not row["checked"]:
            counts["unchecked"] += 1
        elif value is SKIP:
            counts["checked_without_value"] += 1
        else:
            counts["checked"] += 1
            if same(row["field"], imported[(row["id"], row["field"])], value):
                counts["agree"] += 1
            else:
                differ.append({**row, "imported": imported[(row["id"], row["field"])], "value": value})
    for counts in fields.values():
        counts["accuracy"] = round(counts["agree"] / counts["checked"], 3) if counts["checked"] else None
    return {"fields": {field: counts for field, counts in fields.items() if counts["rows"]}, "differ": differ}


def read_filled(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    sheet = load_workbook(path, read_only=True, data_only=True)[f"抽样{SAMPLE_SIZE}道"]
    values = list(sheet.iter_rows(values_only=True))
    col = {name: index for index, name in enumerate(values[0])}
    return [
        {
            "id": cells[col["编号"]],
            "title": cells[col["菜名"]],
            "field": cells[col["字段"]],
            "value": cells[col[YOUR_VALUE]],
            "checked": str(cells[col[CHECKED]] or "").strip() == "是",
        }
        for cells in values[1:]
        if cells[col["编号"]]
    ]


def markdown(result: dict, path: Path, sample: dict) -> str:
    lines = [
        f"# Sample accuracy of {sample['batch']}",
        "",
        f"Filled sample: SHA-256 `{hashlib.sha256(path.read_bytes()).hexdigest()}`; drawn: `sample.json` "
        f"({sample['size']} recipes, seed {sample['seed']}). Only rows marked 是 (checked in person) count.",
        "",
        "| Field | Rows | Checked | Agree | Accuracy | Unchecked | Checked, no value |",
        "| --- | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    for field, c in result["fields"].items():
        share = "n/a" if c["accuracy"] is None else f"{c['accuracy']:.1%}"
        lines.append(
            f"| `{field}` | {c['rows']} | {c['checked']} | {c['agree']} | {share} | {c['unchecked']} "
            f"| {c['checked_without_value']} |"
        )
    lines += ["", f"Disagreements: {len(result['differ'])}.", ""]
    lines += [
        f"- {row['id']} {row['title']}: `{row['field']}` imported `{shown(row['imported'])}`, "
        f"reviewer `{shown(row['value'])}`"
        for row in result["differ"]
    ]
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="Per-field agreement of the reviewer's filled sample.")
    parser.add_argument("filled", type=Path)
    parser.add_argument("--out", type=Path, help="also write the report here (markdown) and its JSON beside it")
    args = parser.parse_args()
    sample = json.loads(SAMPLE.read_text(encoding="utf-8"))
    result = accuracy(read_filled(args.filled), sample)
    report = markdown(result, args.filled, sample)
    print(report)
    if args.out:
        args.out.write_text(report, encoding="utf-8", newline="\n")
        args.out.with_suffix(".json").write_text(
            json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n"
        )


if __name__ == "__main__":
    main()
