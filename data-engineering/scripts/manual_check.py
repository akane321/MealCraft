"""The in-person 100-recipe manual check of a release (03 plan 2.7; owner 2026-10-25): a fixed-seed sample stratified
by course, and the reviewer's three checks on every recipe: complete, steps sensible, servings normal.

    python -I -X utf8 data-engineering/scripts/manual_check.py draw               # sample.json and check.xlsx
    python -I -X utf8 data-engineering/scripts/manual_check.py score <filled.xlsx>    # report.md and report.json

Nothing in the sheet is prefilled. Only rows the reviewer marked 是 count. A recipe is correct when all three checks
are 对 and wrong when any is 错; each 错 counts under its check, so a recipe can count under more than one. A recipe
with no 错 and some 看不出 is reported apart and counts neither way. The correct rate carries a Wilson 95% interval;
the 03 plan's target is 95%.
Offline, with a Python that has openpyxl, from the repository root.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from review_sampling import (  # noqa: E402
    CHECKED,
    REPOSITORY,
    allocate,
    checked,
    interval,
    percent,
    release_recipes,
    wilson,
)

OUT = REPOSITORY / "data-engineering/docs/manual-check-20261025"
SEED = 20261025
SIZE = 100
MINIMUM = 3
TARGET = 0.95
SHEET = "抽检"
CHECKS = ("菜谱完整（对/错）", "步骤合理（对/错）", "份量正常（对/错）")
VERDICTS = ("对", "错", "看不出")
HEAD = (
    "序号",
    "编号",
    "菜名",
    "类别",
    "份数",
    "时间（准备/烹饪/等待 分钟）",
    "每份营养",
    "食材（原文 → 标准食材 克数）",
    "步骤",
    *CHECKS,
    "问题说明",
    CHECKED,
)


def draw(version: str, recipes: list[dict]) -> dict:
    by_course = defaultdict(list)
    for recipe in recipes:
        by_course[recipe["course"]].append(recipe)
    plan = allocate({course: len(members) for course, members in by_course.items()}, SIZE, MINIMUM)
    rng = random.Random(SEED)
    picked = []
    for course, n in sorted(plan.items()):
        picked += rng.sample(sorted(by_course[course], key=lambda r: r["recipe_id"]), n)
    return {
        "seed": SEED,
        "size": len(picked),
        "release": version,
        "by_course": dict(sorted(plan.items())),
        "recipes": [
            {"id": r["recipe_id"], "title": r["title"], "course": r["course"], "servings": r["servings"]}
            for r in sorted(picked, key=lambda r: r["recipe_id"])
        ],
    }


def ingredient_lines(recipe: dict) -> str:
    out = []
    for line in recipe["ingredients"]:
        grams = f"{line['grams']:g} g" if line.get("grams") is not None else "? g"
        optional = "（可选）" if line.get("optional") else ""
        out.append(f"{line['original_text']} → {line['canonical_name']} {grams}{optional}")
    return "\n".join(out)


def workbook(sample: dict, recipes: dict[str, dict]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    notes = book.active
    notes.title = "说明"
    for text in (
        "100 道人工抽检（03 方案 2.7）",
        "",
        f"这 {sample['size']} 道菜按固定种子 {sample['seed']} 从 release 里按类别分层抽出（每类至少 {MINIMUM} 道，"
        "其余按大小分）。",
        "每一行是一道菜。请你本人逐道先看“食材”和“步骤”，然后填黄色几栏。过敏原另有一张表，这里不用看。",
        "",
        "1. 菜谱完整：做这道菜需要的食材都列了吗（步骤里用到、清单里没有的，例如馅料、炸油、高汤，"
        "就是不完整）。填 对 或 错。",
        "2. 步骤合理：照着步骤能做出这道菜吗（步骤缺失、顺序乱、和食材对不上就是不合理）。填 对 或 错。",
        "3. 份量正常：份数和每份的量、热量是否说得过去（例如 1 份 3,000 kcal、"
        "或者 30 份的小饼干每份 2 kcal 就不正常）。"
        "填 对 或 错。",
        "   三项任何一项看不出来，填“看不出”。填“错”或“看不出”时，在“问题说明”里写一句为什么。",
        "4. 已亲自核对：只有你亲自看完这一道菜之后，才填“是”。没看的不要填。没填“是”的行不计入任何结果。",
        "",
        "不要改其他栏，不要增删行。可以用 AI 帮你查不认识的食材，但每一项判断都要你本人确认。",
    ):
        notes.append([text])
    notes.column_dimensions["A"].width = 130
    sheet = book.create_sheet(SHEET)
    sheet.append(list(HEAD))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    yours = PatternFill("solid", fgColor="FFF2CC")
    verdict = DataValidation(type="list", formula1=f'"{",".join(VERDICTS)}"', allow_blank=True)
    only_yes = DataValidation(type="list", formula1='"是"', allow_blank=True)
    sheet.add_data_validation(verdict)
    sheet.add_data_validation(only_yes)
    for number, drawn in enumerate(sample["recipes"], start=1):
        recipe = recipes[drawn["id"]]
        nutrition = recipe.get("nutrition") or {}
        per_serving = (
            f"{nutrition.get('energy_kcal')} kcal, 蛋白 {nutrition.get('protein_g')} g, "
            f"碳水 {nutrition.get('carbohydrate_g')} g, 脂肪 {nutrition.get('fat_g')} g"
        )
        times = f"{recipe.get('prep_minutes')}/{recipe.get('cook_minutes')}/{recipe.get('passive_minutes')}"
        steps = "\n".join(f"{step['step_number']}. {step['text']}" for step in recipe["instructions"])
        sheet.append(
            [
                number,
                drawn["id"],
                recipe["title"],
                recipe["course"],
                recipe["servings"],
                times,
                per_serving,
                ingredient_lines(recipe),
                steps,
                None,
                None,
                None,
                None,
                None,
            ]
        )
        row = sheet.max_row
        for column in "JKLMN":
            sheet[f"{column}{row}"].fill = yours
        for column in "JKL":
            verdict.add(f"{column}{row}")
        only_yes.add(f"N{row}")
    widths = (5, 22, 28, 12, 6, 12, 22, 60, 70, 10, 10, 10, 30, 10)
    for column, width in zip("ABCDEFGHIJKLMN", widths, strict=True):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "D2"
    return book


def read_filled(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    values = list(load_workbook(path, read_only=True, data_only=True)[SHEET].iter_rows(values_only=True))
    column = {name: index for index, name in enumerate(values[0])}
    return [
        {
            "id": cells[column["编号"]],
            "title": cells[column["菜名"]],
            "verdicts": [str(cells[column[check]] or "").strip() for check in CHECKS],
            "note": cells[column["问题说明"]],
            "checked": checked(cells[column[CHECKED]]),
        }
        for cells in values[1:]
        if cells[column["编号"]]
    ]


def score(rows: list[dict], sample: dict) -> dict:
    drawn = {recipe["id"]: recipe for recipe in sample["recipes"]}
    if sorted(row["id"] for row in rows) != sorted(drawn):
        raise SystemExit(
            f"the filled sheet's rows are not the sample's: {sorted({r['id'] for r in rows} ^ set(drawn))}"
        )
    names = [check.split("（")[0] for check in CHECKS]
    wrong, by_course = Counter(), defaultdict(lambda: {"checked": 0, "correct": 0})
    problems, unsure, refill = [], [], []
    unchecked = 0
    for row in rows:
        if not row["checked"]:
            unchecked += 1
            continue
        if any(verdict not in VERDICTS for verdict in row["verdicts"]):
            refill.append({"id": row["id"], "title": row["title"], "verdicts": row["verdicts"]})
            continue
        if "看不出" in row["verdicts"] and "错" not in row["verdicts"]:
            unsure.append({"id": row["id"], "title": row["title"], "note": row["note"]})
            continue
        course = by_course[drawn[row["id"]]["course"]]
        course["checked"] += 1
        failed = [name for name, verdict in zip(names, row["verdicts"], strict=True) if verdict == "错"]
        course["correct"] += not failed
        wrong.update(failed)
        if failed:
            problems.append({"id": row["id"], "title": row["title"], "failed": failed, "note": row["note"]})
    total = sum(course["checked"] for course in by_course.values())
    correct = sum(course["correct"] for course in by_course.values())
    errors = sum(wrong.values())
    return {
        "sampled": len(rows),
        "checked": total,
        "correct": correct,
        "rate": round(correct / total, 3) if total else None,
        "wilson95": wilson(correct, total),
        "target": TARGET,
        "unsure": unsure,
        "unchecked": unchecked,
        "refill": refill,
        "errors_by_check": {
            name: {"recipes": wrong[name], "share_of_errors": round(wrong[name] / errors, 3) if errors else None}
            for name in names
        },
        "by_course": {name: by_course[name] for name in sorted(by_course)},
        "problems": problems,
    }


def markdown(result: dict, sample: dict, digest: str) -> str:
    rate = result["rate"]
    met = "n/a" if rate is None else ("met" if rate >= result["target"] else "not met")
    lines = [
        f"# Manual check of {sample['size']} recipes of release {sample['release']}",
        "",
        "Generated by `data-engineering/scripts/manual_check.py score`; do not edit. "
        f"Filled sheet SHA-256 `{digest}`; sample `sample.json` (seed {sample['seed']}, stratified by course). "
        "Only rows the reviewer marked 是 count.",
        "",
        f"- Sampled {result['sampled']}, checked in person {result['checked']}, unchecked {result['unchecked']}, "
        f"with a 看不出 {len(result['unsure'])}, to refill {len(result['refill'])}.",
        f"- Correct on all three checks: {result['correct']} of {result['checked']}, **{percent(rate)}** "
        f"(Wilson 95% {interval(result['wilson95'])}); target {percent(result['target'])}: {met}.",
        "",
        "## Errors by check",
        "",
        "A recipe that fails two checks counts under both.",
        "",
        "| Check | Recipes | Share of errors |",
        "| --- | ---: | ---: |",
    ]
    for name, counts in result["errors_by_check"].items():
        lines.append(f"| {name} | {counts['recipes']} | {percent(counts['share_of_errors'])} |")
    lines += ["", "## By course", "", "| Course | Checked | Correct |", "| --- | ---: | ---: |"]
    for name, counts in result["by_course"].items():
        lines.append(f"| `{name}` | {counts['checked']} | {counts['correct']} |")
    lines += ["", f"## Problems ({len(result['problems'])})", ""]
    lines += [
        f"- {row['id']} {row['title']}: {', '.join(row['failed'])}" + (f" ({row['note']})" if row["note"] else "")
        for row in result["problems"]
    ]
    if result["unsure"]:
        lines += ["", "## Could not tell", ""]
        lines += [
            f"- {row['id']} {row['title']}" + (f" ({row['note']})" if row["note"] else "") for row in result["unsure"]
        ]
    if result["refill"]:
        lines += ["", "## To refill", ""]
        lines += [f"- {row['id']} {row['title']}: {row['verdicts']}" for row in result["refill"]]
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="The in-person 100-recipe manual check of a release (03 plan 2.7).")
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("draw", help="draw sample.json and write check.xlsx")
    scoring = commands.add_parser("score", help="score the filled check.xlsx into report.md and report.json")
    scoring.add_argument("filled", type=Path)
    args = parser.parse_args()
    if args.command == "draw":
        version, recipes = release_recipes()
        sample = draw(version, recipes)
        (OUT / "sample.json").write_text(json.dumps(sample, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
        workbook(sample, {recipe["recipe_id"]: recipe for recipe in recipes}).save(OUT / "check.xlsx")
        print(json.dumps(sample["by_course"], ensure_ascii=False))
        return
    sample = json.loads((OUT / "sample.json").read_text(encoding="utf-8"))
    result = score(read_filled(args.filled), sample)
    report = markdown(result, sample, hashlib.sha256(args.filled.read_bytes()).hexdigest())
    (OUT / "report.md").write_text(report, encoding="utf-8", newline="\n")
    (OUT / "report.json").write_text(
        json.dumps(result, indent=1, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n"
    )
    print(report)


if __name__ == "__main__":
    main()
