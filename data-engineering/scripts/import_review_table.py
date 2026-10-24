"""Import the human reviewer's workbook for the frozen model batch as an AI-estimated batch (owner 2026-10-24).

The workbook's decisions were prefilled by AI and confirmed in chat by dish type, not checked row by row: its
"已亲自核对" column is empty on every row. So nothing here is a human label. ADR-0076's route:

- only a field the current rules leave open (`rule_suggestion` null in the reviewer packet) is imported, into
  `recipe-tags-v1/review-table-20261024` as model-decided and unreviewed: a role the reviewer may choose
  (`review_allowed_roles`), its staple (true exactly for one_bowl and staple), and protein sources of the vocabulary;
- a field the rules or the owner decide, a null, a skipped row, a role outside the reviewer's options, a staple that
  disagrees with its role, or an unknown code is not imported, and is listed with its reason;
- the reviewer checks a fixed-seed sample of 30 of the recipes in person (`sample.xlsx`) and
  `review_table_accuracy.py` reports the sample's accuracy, which is published. Only then may the owner accept the
  batch (`python -m app.data.recipe_tags --accept-batch recipe-tags-v1/review-table-20261024`).

Offline, with a Python that has openpyxl, from the repository root:
`python -I -X utf8 data-engineering/scripts/import_review_table.py <workbook.xlsx> --base <commit>`.
It writes `decisions.json`, `sample.json` and `sample.xlsx` and never the tag file: `python -m app.data.recipe_tags`
reads `decisions.json`.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app.data.recipe_tags import (  # noqa: E402
    MEAL_ROLES,
    OWNER_ROLES,
    PROTEIN_SOURCES,
    REVIEW_TABLE,
    REVIEW_TABLE_BATCH,
)
from app.data.release_v2 import release_dir  # noqa: E402

PACKET = "data-engineering/docs/role-rules-20261023b/review.json"
OUT = ROOT / REVIEW_TABLE
SHEETS = ("1 可用池277道", "2 其余145道")
FIELDS = ("meal_role", "contains_staple", "protein_sources")
STAPLE_ROLES = frozenset({"one_bowl", "staple"})
SKIP = "<skip>"  # an empty cell or decision, or one left to the owner: the reviewer gave no value
PENDING = "待组长定"
SEED = 20261024
SAMPLE_SIZE = 30
STEPS_SHOWN = 400
CHECKED = "已亲自核对（填 是）"
YOUR_VALUE = "你看完后的值"
REASONS = (
    "rule_decided_agrees",
    "rule_decided_conflicts",
    "owner_decided",
    "outside_allowed_roles",
    "left_open",
    "skipped_by_reviewer",
    "staple_inconsistent",
    "invalid_value",
)


def code(text):
    """A workbook cell as a value: 'staple（主食）' -> 'staple', 'pork（猪）,beef（牛）' -> ['beef', 'pork'],
    'true（是）' -> True, '[]' -> [], 'null（未知/不放行）' -> None, an empty cell or '待组长定' -> SKIP.
    A leading '被规则挡住 →' is dropped; each comma-separated part is cut at its first bracket. Excel's TRUE and
    FALSE cells are booleans already."""
    if isinstance(text, bool):
        return text
    if text is None or not str(text).strip():
        return SKIP
    text = re.sub(r"^被规则挡住\s*→\s*", "", str(text).strip().lower())
    text = ",".join(re.split(r"[（(]", part)[0].strip() for part in text.split(","))
    if text == PENDING:
        return SKIP
    if text in {"null", ""}:
        return None
    if text in {"true", "false"}:
        return text == "true"
    if text == "[]":
        return []
    if "," in text or text in PROTEIN_SOURCES:
        return sorted(part.strip() for part in text.split(",") if part.strip())
    return text


def intended(row: dict):
    """The value a workbook row's decision means: Claude's suggestion, the model's value, or the other value; a chat
    confirmation or a blocked row with no other value means the suggestion. No decision, or 跳过, is SKIP."""
    decision = str(row["decision"] or "").strip()
    if not decision or decision.startswith("跳过"):
        return SKIP
    if decision.startswith("同意Claude建议"):
        return code(row["claude"])
    if decision.startswith("同意模型原值"):
        return code(row["model"])
    if row["other"] not in (None, ""):
        return code(row["other"])
    if decision.startswith(("聊天确认", "被规则挡住")):
        return code(row["claude"])
    return SKIP


def judge(key: str, field: str, value, current: dict) -> str | None:
    """Why a workbook value is not imported, or None when it may be (a staple is then checked against its role)."""
    if value is SKIP:
        return "skipped_by_reviewer"
    if field != "protein_sources" and key in OWNER_ROLES:
        return "owner_decided"
    rule = current["rule_suggestion"][field]
    if rule is not None:
        agrees = sorted(rule) == value if field == "protein_sources" else rule == value and type(rule) is type(value)
        return "rule_decided_agrees" if agrees else "rule_decided_conflicts"
    if value is None:
        return "left_open"
    if field == "meal_role":
        if value not in MEAL_ROLES:
            return "invalid_value"
        return None if value in current["review_allowed_roles"] else "outside_allowed_roles"
    if field == "contains_staple":
        return None if isinstance(value, bool) else "invalid_value"
    return None if isinstance(value, list) and set(value) <= set(PROTEIN_SOURCES) else "invalid_value"


def decide(rows: list[dict], packet: dict[str, dict]) -> tuple[dict, list[dict]]:
    """The imported values by recipe, and every workbook value not imported with its reason."""
    by_recipe: dict[str, dict] = defaultdict(dict)
    for row in rows:
        by_recipe[row["id"]][row["field"]] = row
    missing = sorted(set(by_recipe) - set(packet))
    if missing:
        raise SystemExit(f"recipes not in the reviewer packet: {missing[:5]}")
    accepted, skipped = {}, []
    for key in sorted(by_recipe):
        current = packet[key]
        values = {field: intended(row) for field, row in by_recipe[key].items()}
        reasons = {field: judge(key, field, value, current) for field, value in values.items()}
        if "meal_role" in values and reasons["meal_role"] is None:
            derived = values["meal_role"] in STAPLE_ROLES
            staple = values.get("contains_staple", SKIP)
            if staple is not SKIP and staple is not derived:
                reasons["meal_role"] = reasons["contains_staple"] = "staple_inconsistent"
            else:
                values["contains_staple"], reasons["contains_staple"] = derived, None
        elif reasons.get("contains_staple", "") is None:
            reasons["contains_staple"] = reasons.get("meal_role") or "left_open"  # a staple follows its role
        kept = {field: values[field] for field in FIELDS if field in reasons and reasons[field] is None}
        if "protein_sources" in kept:
            kept["protein_sources"] = [source for source in PROTEIN_SOURCES if source in kept["protein_sources"]]
        if kept:
            accepted[key] = {"title": current["title"], **kept}
        skipped += [
            {
                "id": key,
                "title": current["title"],
                "field": field,
                "value": None if values[field] is SKIP else values[field],
                "rule": current["rule_suggestion"][field],
                "reason": reason,
            }
            for field, reason in reasons.items()
            if reason
        ]
    return accepted, skipped


def counts(accepted: dict, skipped: list[dict]) -> dict:
    by_field = {field: Counter() for field in FIELDS}
    for values in accepted.values():
        for field in FIELDS:
            by_field[field]["imported"] += field in values
    for item in skipped:
        by_field[item["field"]][item["reason"]] += 1
    reasons = Counter(item["reason"] for item in skipped)
    return {
        "recipes_with_an_imported_value": len(accepted),
        "imported": {field: by_field[field]["imported"] for field in FIELDS},
        "skipped": {reason: reasons[reason] for reason in REASONS},
        "by_field": {field: {name: n for name, n in by_field[field].items() if n} for field in FIELDS},
        "imported_meal_roles": dict(
            sorted(Counter(v["meal_role"] for v in accepted.values() if "meal_role" in v).items())
        ),
    }


def draw_sample(accepted: dict) -> list[tuple[str, str]]:
    """`SAMPLE_SIZE` recipes with an imported value, stratified by the imported meal role (protein sources alone are
    their own stratum): every stratum at least one, the rest by size (largest remainder), drawn with `SEED`."""
    strata: dict[str, list[str]] = defaultdict(list)
    for key in sorted(accepted):
        strata[accepted[key].get("meal_role", "protein_sources only")].append(key)
    total = len(accepted)
    quota = {name: max(1, len(keys) * SAMPLE_SIZE // total) for name, keys in strata.items()}
    while sum(quota.values()) < min(SAMPLE_SIZE, total):
        name = max(sorted(strata), key=lambda s: len(strata[s]) * SAMPLE_SIZE / total - quota[s])
        quota[name] += 1
    rng = random.Random(SEED)
    return [(name, key) for name in sorted(strata) for key in sorted(rng.sample(strata[name], quota[name]))]


def shown(value) -> str:
    """A value as the code the reviewer reads and types."""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, list):
        return ",".join(value) if value else "[]"
    return "null" if value is None else str(value)


def read_rows(path: Path) -> tuple[list[dict], int]:
    """The workbook's field rows, and how many are marked personally checked."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in SHEETS:
        values = list(workbook[sheet].iter_rows(values_only=True))
        col = {name: index for index, name in enumerate(values[0])}
        other = col[next(name for name in col if str(name).startswith("其他值"))]
        checked = col[next(name for name in col if str(name).startswith("已亲自核对"))]
        for cells in values[1:]:
            if cells[col["编号"]]:
                rows.append(
                    {
                        "id": cells[col["编号"]],
                        "field": re.split(r"[（(]", str(cells[col["字段"]]))[0],
                        "model": cells[col["AI模型原值"]],
                        "claude": cells[col["Claude建议"]],
                        "decision": cells[col["你的决定"]],
                        "other": cells[other],
                        "checked": str(cells[checked] or "").strip() == "是",
                    }
                )
    return rows, sum(row["checked"] for row in rows)


def steps_of(keys: set[str]) -> dict[str, str]:
    steps = {}
    with (release_dir() / "recipes.jsonl").open(encoding="utf-8") as handle:
        for record in map(json.loads, filter(str.strip, handle)):
            if record["recipe_id"] in keys:
                text = " / ".join(s["text"].strip() for s in record["instructions"] if s["text"].strip())
                steps[record["recipe_id"]] = text[:STEPS_SHOWN]
    return steps


INSTRUCTIONS = [
    "AI 标签抽样核对（recipe-tags-v1/review-table-20261024）",
    "",
    f"这 {SAMPLE_SIZE} 道菜按固定种子 {SEED} 抽自“复核表里的值被导入的菜”，按导入的餐型角色分层。",
    "导入值来自你的复核表，是 AI 预填、在聊天里按菜品类型确认的，没有人逐条看过食材，所以现在仍记为 AI 估计、未复核。",
    "这次核对的结果决定这一批能不能被组长接受。",
    "",
    "每一行是一道菜的一个字段。请你本人逐行：",
    "1. 先看“食材（克）”和“步骤（前 400 字）”，再看“导入值”。不要只凭菜名判断。",
    "2. 在“你看完后的值”里填你认为正确的值（同意导入值也照抄一遍，不能空着），用代码：",
    "   餐型角色：" + " / ".join(MEAL_ROLES) + "，并且必须在“可选角色”里；",
    "   是否自带主食：true 或 false（one_bowl 和 staple 是 true，其他角色是 false）；",
    "   主要蛋白来源：用英文逗号分隔 " + ", ".join(PROTEIN_SOURCES) + "；没有主要蛋白填 []；",
    "   看不出来、或正确的值不在可选范围里：填 null。",
    "3. 只有你亲自看完这一行的食材和步骤之后，才在“已亲自核对（填 是）”里填“是”。没看的行不要填。",
    "   没填“是”的行不计入准确率，只计为“未核对”。",
    "4. 不要改其他栏，不要增删行；有话可写在“备注”里。",
    "",
    "准确率由 data-engineering/scripts/review_table_accuracy.py 计算并公开。组长看过准确率后，",
    "才决定是否接受这一批（python -m app.data.recipe_tags --accept-batch recipe-tags-v1/review-table-20261024）。",
]
SAMPLE_COLUMNS = (
    "序号",
    "编号",
    "菜名",
    "类别",
    "食材（克）",
    "步骤（前 400 字）",
    "字段",
    "导入值",
    "可选角色",
    YOUR_VALUE,
    CHECKED,
    "备注",
)


def write_sample_xlsx(path: Path, recipes: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    guide = workbook.active
    guide.title = "说明"
    for line in INSTRUCTIONS:
        guide.append([line])
    guide["A1"].font = Font(bold=True)
    guide.column_dimensions["A"].width = 110
    sheet = workbook.create_sheet(f"抽样{SAMPLE_SIZE}道")
    sheet.append(SAMPLE_COLUMNS)
    for number, recipe in enumerate(recipes, 1):
        ingredients = "; ".join(
            f"{line['name']} {line['grams'] if line['grams'] is not None else '?'} g（{line['text']}）"
            for line in recipe["ingredients"]
        )
        for field in FIELDS:
            if field in recipe["imported"]:
                roles = " / ".join(recipe["review_allowed_roles"]) if field == "meal_role" else ""
                sheet.append(
                    [number, recipe["id"], recipe["title"], recipe["course"], ingredients, recipe["steps"], field]
                    + [shown(recipe["imported"][field]), roles, None, None, None]
                )
    widths = (6, 20, 30, 11, 60, 60, 16, 16, 24, 16, 14, 24)
    fill = PatternFill("solid", fgColor="FFF2CC")
    for index, width in enumerate(widths):
        letter = chr(ord("A") + index)
        sheet.column_dimensions[letter].width = width
        sheet[f"{letter}1"].font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[9:12]:
            cell.fill = fill
    sheet.freeze_panes = "C2"
    workbook.save(path)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description="Import the reviewer's workbook as an unreviewed AI-estimated batch.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--base", help="the commit the import is made against (default: HEAD)")
    args = parser.parse_args()
    base = args.base or subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    packet_path = ROOT / PACKET
    packet = {row["id"]: row for row in json.loads(packet_path.read_text(encoding="utf-8"))["frozen_rows"]}
    rows, checked = read_rows(args.workbook)
    accepted, skipped = decide(rows, packet)
    summary = counts(accepted, skipped)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    decisions = {
        "batch": REVIEW_TABLE_BATCH,
        "method": "llm",
        "reviewed": False,
        "note": (
            "AI-prefilled values the reviewer confirmed in chat by dish type, not checked row by row. Imported as an "
            "unreviewed model-decided batch, only where the current rules leave the field open; never a human label."
        ),
        "provenance": {
            "workbook": "the reviewer's review workbook (not committed)",
            "workbook_sha256": sha256(args.workbook),
            "sheets": list(SHEETS),
            "rows": len(rows),
            "recipes": len({row["id"] for row in rows}),
            "personally_checked_rows": checked,
            "personally_checked_note": "the 已亲自核对 column is empty on every row" if not checked else "not empty",
            "packet": PACKET,
            "packet_sha256": sha256(packet_path),
            "base_commit": base,
            "owner_decisions": "backend/app/data/recipe_tags.py OWNER_ROLES (owner 2026-10-24) win over the workbook",
            "generated_by": "data-engineering/scripts/import_review_table.py",
        },
        "counts": summary,
        "accepted": accepted,
        "skipped": skipped,
    }
    OUT.write_text(json.dumps(decisions, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")

    picked = draw_sample(accepted)
    steps = steps_of({key for _, key in picked})
    recipes = [
        {
            "id": key,
            "title": packet[key]["title"],
            "course": packet[key]["course"],
            "stratum": stratum,
            "imported": {field: accepted[key][field] for field in FIELDS if field in accepted[key]},
            "review_allowed_roles": packet[key]["review_allowed_roles"],
            "ingredients": packet[key]["ingredients"],
            "steps": steps.get(key, ""),
        }
        for stratum, key in picked
    ]
    strata = Counter(accepted[key].get("meal_role", "protein_sources only") for key in accepted)
    sample = {
        "batch": REVIEW_TABLE_BATCH,
        "seed": SEED,
        "size": len(recipes),
        "drawn_from": len(accepted),
        "strata": {
            name: {"recipes": n, "drawn": sum(s == name for s, _ in picked)} for name, n in sorted(strata.items())
        },
        "how_to_check": "data-engineering/docs/review-table-20261024/README.md; fill sample.xlsx in person",
        "recipes": [
            {name: recipe[name] for name in ("id", "title", "course", "stratum", "imported")} for recipe in recipes
        ],
    }
    (OUT.parent / "sample.json").write_text(
        json.dumps(sample, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n"
    )
    write_sample_xlsx(OUT.parent / "sample.xlsx", recipes)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"sample: {len(recipes)} recipes; strata {sample['strata']}")


if __name__ == "__main__":
    main()
