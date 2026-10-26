"""The in-person allergen review of a release (03 plan 2.5; owner 2026-10-25): a fixed-seed sample of 50 recipes,
stratified by allergen kind, AI candidates beside the catalog's allergens, and the reviewer's confirmed allergens.

    python -I -X utf8 data-engineering/scripts/allergen_review.py draw              # sample.json and review.xlsx
    python -I -X utf8 data-engineering/scripts/allergen_review.py score <filled.xlsx>   # report.md and report.json

A recipe's stratum is the rarest allergen the catalog gives it, or "none" when it has none, so every allergen kind is
seen at least MINIMUM times and the misses on "no allergen" recipes are measured too. Only rows the reviewer marked
是 count. A recipe agrees when the confirmed set equals the catalog's; each stratum carries a Wilson 95% interval and
the overall agreement weights each stratum by its share of the release, since rare allergens are over-sampled.
`ai_candidates.json` holds Claude's candidates, read from the ingredient text and steps only; they are scored against
the reviewer the same way and never stand in for the reviewer. Offline, with a Python that has openpyxl, from the
repository root.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from review_sampling import (  # noqa: E402
    CHECKED,
    REPOSITORY,
    allocate,
    checked,
    interval,
    percent,
    release_recipes,
    wilson,
)

OUT = REPOSITORY / "data-engineering/docs/allergen-review-20261025"
SEED = 20261025
SIZE = 50
MINIMUM = 3
SHEET = "复核"
CONFIRMED = "你确认的过敏原"
CODES = (
    "gluten",
    "gluten_candidate",
    "milk",
    "eggs",
    "soy",
    "fish",
    "crustaceans",
    "molluscs",
    "peanuts",
    "tree_nuts",
    "sesame",
    "sulfites",
)
HEAD = (
    "序号",
    "编号",
    "菜名",
    "类别",
    "抽样层",
    "食材（原文 → 标准食材 [目录给该食材的过敏原]）",
    "步骤",
    "目录里的过敏原",
    "AI 候选",
    "AI 说明",
    CONFIRMED,
    "备注",
    CHECKED,
)


def draw(version: str, recipes: list[dict]) -> dict:
    frequency = Counter(allergen for recipe in recipes for allergen in recipe["allergens"])
    strata = defaultdict(list)
    for recipe in recipes:
        allergens = recipe["allergens"]
        strata[min(allergens, key=lambda a: (frequency[a], a)) if allergens else "none"].append(recipe)
    plan = allocate({name: len(members) for name, members in sorted(strata.items())}, SIZE, MINIMUM)
    rng = random.Random(SEED)
    picked = []
    for name in sorted(plan):
        for recipe in rng.sample(sorted(strata[name], key=lambda r: r["recipe_id"]), plan[name]):
            picked.append(
                {
                    "id": recipe["recipe_id"],
                    "stratum": name,
                    "stratum_size": len(strata[name]),
                    "allergens": sorted(recipe["allergens"]),
                }
            )
    return {
        "seed": SEED,
        "size": len(picked),
        "release": version,
        "recipes_in_release": len(recipes),
        "strata": {name: {"recipes": len(strata[name]), "drawn": plan[name]} for name in sorted(plan)},
        "recipes": sorted(picked, key=lambda r: r["id"]),
    }


def shown(values) -> str:
    return ", ".join(values) if values else "[]"


def workbook(sample: dict, ai: dict, recipes: dict[str, dict]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    notes = book.active
    notes.title = "说明"
    for text in (
        "过敏原复核（03 方案 2.5）",
        "",
        f"这 {sample['size']} 道菜按固定种子 {sample['seed']} 从 release {sample['release']} 里抽出，按过敏原种类分层："
        f"每道菜归到它目录里最少见的那种过敏原（没有过敏原的单独一层），每层至少 {MINIMUM} 道，其余按层的大小分。",
        "“目录里的过敏原”是现在规则算出来的值，要复核的就是它。“AI 候选”是 Claude 只看食材原文和步骤给的判断，"
        "和目录不同时写了理由。",
        "",
        "请你本人逐道：",
        "1. 先看“食材”和“步骤”，再看“目录里的过敏原”和“AI 候选”。",
        "2. 在“你确认的过敏原”里填你认为正确的全部过敏原（同意也照抄一遍，不能空着），用英文代码、逗号分隔：",
        "   " + ", ".join(CODES),
        "   gluten_candidate = 可能含麸质、要看具体商品（例如没写明的 tortilla、米酒）。"
        "一个都没有填 []。实在看不出填 null。",
        "3. 你认为目录错了的，在“备注”里写一句为什么（例如“flour tortilla 一定含麸质”）。",
        "4. 只有亲自看完这一道菜之后，才在“已亲自核对”填“是”。没填“是”的不计入结果。",
        "",
        "口径（和目录一致）：配菜建议（serve with ...）和可选的（if desired）不算；"
        "“A 或 B”两样都可能用，两样的过敏原都算。",
        "不要改其他栏，不要增删行。",
    ):
        notes.append([text])
    notes.column_dimensions["A"].width = 130
    sheet = book.create_sheet(SHEET)
    sheet.append(list(HEAD))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    yours = PatternFill("solid", fgColor="FFF2CC")
    only_yes = DataValidation(type="list", formula1='"是"', allow_blank=True)
    sheet.add_data_validation(only_yes)
    for number, drawn in enumerate(sample["recipes"], start=1):
        recipe = recipes[drawn["id"]]
        lines = "\n".join(
            f"{line['original_text']} → {line['canonical_name']}"
            + (f" [{', '.join(line['allergens'])}]" if line["allergens"] else "")
            for line in recipe["ingredients"]
        )
        steps = "\n".join(f"{step['step_number']}. {step['text']}" for step in recipe["instructions"])
        candidate = ai[drawn["id"]]
        sheet.append(
            [
                number,
                drawn["id"],
                recipe["title"],
                recipe["course"],
                drawn["stratum"],
                lines,
                steps,
                shown(drawn["allergens"]),
                shown(candidate["allergens"]),
                candidate["note"],
                None,
                None,
                None,
            ]
        )
        row = sheet.max_row
        for column in "KLM":
            sheet[f"{column}{row}"].fill = yours
        only_yes.add(f"M{row}")
    for column, width in zip("ABCDEFGHIJKLM", (5, 22, 28, 12, 14, 60, 70, 22, 22, 40, 22, 30, 10), strict=True):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "D2"
    return book


def parse(text) -> set[str] | str | None:
    """The confirmed allergens: a set, None for null (cannot tell), or a string saying why the cell must be refilled."""
    if text is None or not str(text).strip():
        return "empty"
    text = str(text).strip().lower()
    if text == "null":
        return None
    if text == "[]":
        return set()
    parts = {part.strip() for part in re.split(r"[,，、\s]+", text) if part.strip()}
    unknown = parts - set(CODES)
    return f"unknown code {sorted(unknown)}" if unknown else parts


def read_filled(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    values = list(load_workbook(path, read_only=True, data_only=True)[SHEET].iter_rows(values_only=True))
    column = {name: index for index, name in enumerate(values[0])}
    return [
        {
            "id": cells[column["编号"]],
            "title": cells[column["菜名"]],
            "confirmed": cells[column[CONFIRMED]],
            "note": cells[column["备注"]],
            "checked": checked(cells[column[CHECKED]]),
        }
        for cells in values[1:]
        if cells[column["编号"]]
    ]


def score(rows: list[dict], sample: dict, ai: dict) -> dict:
    drawn = {recipe["id"]: recipe for recipe in sample["recipes"]}
    if sorted(row["id"] for row in rows) != sorted(drawn):
        raise SystemExit(
            f"the filled sheet's rows are not the sample's: {sorted({r['id'] for r in rows} ^ set(drawn))}"
        )
    strata = {name: {"checked": 0, "agree": 0} for name in sample["strata"]}
    missed, extra, corrections, refill = Counter(), Counter(), [], []
    unchecked = cannot_tell = ai_agree = 0
    for row in rows:
        if not row["checked"]:
            unchecked += 1
            continue
        confirmed = parse(row["confirmed"])
        if isinstance(confirmed, str):
            refill.append({"id": row["id"], "title": row["title"], "why": confirmed})
            continue
        if confirmed is None:
            cannot_tell += 1
            continue
        recipe = drawn[row["id"]]
        catalog = set(recipe["allergens"])
        counts = strata[recipe["stratum"]]
        counts["checked"] += 1
        counts["agree"] += confirmed == catalog
        ai_agree += confirmed == set(ai[row["id"]]["allergens"])
        missed.update(confirmed - catalog)
        extra.update(catalog - confirmed)
        if confirmed != catalog:
            corrections.append(
                {
                    "id": row["id"],
                    "title": row["title"],
                    "catalog": sorted(catalog),
                    "confirmed": sorted(confirmed),
                    "missed": sorted(confirmed - catalog),
                    "extra": sorted(catalog - confirmed),
                    "note": row["note"],
                }
            )
    weighted = covered = 0
    for name, counts in strata.items():
        counts["recipes_in_release"] = sample["strata"][name]["recipes"]
        counts["agreement"] = round(counts["agree"] / counts["checked"], 3) if counts["checked"] else None
        counts["wilson95"] = wilson(counts["agree"], counts["checked"])
        if counts["checked"]:
            weighted += counts["agree"] / counts["checked"] * counts["recipes_in_release"]
            covered += counts["recipes_in_release"]
    total_checked = sum(counts["checked"] for counts in strata.values())
    total_agree = sum(counts["agree"] for counts in strata.values())
    return {
        "sampled": len(rows),
        "checked": total_checked,
        "agree": total_agree,
        "unchecked": unchecked,
        "cannot_tell": cannot_tell,
        "refill": refill,
        "agreement_unweighted": round(total_agree / total_checked, 3) if total_checked else None,
        "wilson95_unweighted": wilson(total_agree, total_checked),
        "agreement_weighted": round(weighted / covered, 3) if covered else None,
        "weighted_cover": {"recipes": covered, "of": sample["recipes_in_release"]},
        "strata": strata,
        "missed_by_catalog": dict(sorted(missed.items(), key=lambda item: (-item[1], item[0]))),
        "extra_in_catalog": dict(sorted(extra.items(), key=lambda item: (-item[1], item[0]))),
        "corrections": corrections,
        "ai_candidates": {"agree": ai_agree, "checked": total_checked, "wilson95": wilson(ai_agree, total_checked)},
    }


def markdown(result: dict, sample: dict, digest: str) -> str:
    cover = result["weighted_cover"]
    lines = [
        f"# Allergen review of release {sample['release']}",
        "",
        "Generated by `data-engineering/scripts/allergen_review.py score`; do not edit. "
        f"Filled sheet SHA-256 `{digest}`; sample `sample.json` ({sample['size']} recipes, seed {sample['seed']}, "
        "stratified by the rarest catalog allergen of each recipe). Only rows the reviewer marked 是 count.",
        "",
        f"- Sampled {result['sampled']}, marked 是 (checked) {result['checked']}, unchecked {result['unchecked']}, "
        f"marked null (cannot tell) {result['cannot_tell']}, to refill {len(result['refill'])}.",
        f"- Recipes whose catalog allergens equal the reviewer's: {result['agree']} of {result['checked']} "
        f"({percent(result['agreement_unweighted'])}, Wilson 95% {interval(result['wilson95_unweighted'])}).",
        f"- Weighted by each stratum's share of the release: **{percent(result['agreement_weighted'])}** "
        f"(strata covering {cover['recipes']} of {cover['of']} release recipes).",
        "",
        "| Stratum | Recipes in release | Checked | Agree | Agreement | Wilson 95% |",
        "| --- | ---: | ---: | ---: | ---: | --- |",
    ]
    for name, counts in result["strata"].items():
        lines.append(
            f"| `{name}` | {counts['recipes_in_release']} | {counts['checked']} | {counts['agree']} "
            f"| {percent(counts['agreement'])} | {interval(counts['wilson95'])} |"
        )
    lines += [
        "",
        "## Errors by kind",
        "",
        "Missed: the reviewer found the allergen, the catalog did not. Extra: the catalog has it, the reviewer "
        "did not.",
        "",
        "| Allergen | Missed by the catalog | Extra in the catalog |",
        "| --- | ---: | ---: |",
    ]
    for code in sorted(set(result["missed_by_catalog"]) | set(result["extra_in_catalog"])):
        lines.append(
            f"| `{code}` | {result['missed_by_catalog'].get(code, 0)} | {result['extra_in_catalog'].get(code, 0)} |"
        )
    lines += ["", f"## Corrections ({len(result['corrections'])})", ""]
    lines += [
        f"- {row['id']} {row['title']}: catalog `{shown(row['catalog'])}`, reviewer `{shown(row['confirmed'])}`"
        + (f" ({row['note']})" if row["note"] else "")
        for row in result["corrections"]
    ]
    if result["refill"]:
        lines += ["", "## To refill", ""]
        lines += [f"- {row['id']} {row['title']}: {row['why']}" for row in result["refill"]]
    ai = result["ai_candidates"]
    lines += [
        "",
        "## AI candidates",
        "",
        f"Claude's candidates equal the reviewer's allergens on {ai['agree']} of {ai['checked']} checked recipes "
        f"(Wilson 95% {interval(ai['wilson95'])}). They were shown to the reviewer beside the catalog, so this is not "
        "an independent accuracy of the AI.",
    ]
    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="The in-person allergen review of a release (03 plan 2.5).")
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("draw", help="draw sample.json and write review.xlsx")
    scoring = commands.add_parser("score", help="score the filled review.xlsx into report.md and report.json")
    scoring.add_argument("filled", type=Path)
    args = parser.parse_args()
    if args.command == "draw":
        version, recipes = release_recipes()
        sample = draw(version, recipes)
        (OUT / "sample.json").write_text(json.dumps(sample, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
        ai = json.loads((OUT / "ai_candidates.json").read_text(encoding="utf-8"))
        by_id = {recipe["recipe_id"]: recipe for recipe in recipes}
        workbook(sample, ai, by_id).save(OUT / "review.xlsx")
        print(json.dumps(sample["strata"], ensure_ascii=False))
        return
    sample = json.loads((OUT / "sample.json").read_text(encoding="utf-8"))
    ai = json.loads((OUT / "ai_candidates.json").read_text(encoding="utf-8"))
    result = score(read_filled(args.filled), sample, ai)
    report = markdown(result, sample, hashlib.sha256(args.filled.read_bytes()).hexdigest())
    (OUT / "report.md").write_text(report, encoding="utf-8", newline="\n")
    (OUT / "report.json").write_text(
        json.dumps(result, indent=1, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n"
    )
    print(report)


if __name__ == "__main__":
    main()
